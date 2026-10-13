import os
os.environ.setdefault('MPLCONFIGDIR', '/tmp/trw-matplotlib')
os.environ.setdefault('XDG_CACHE_HOME', '/tmp/trw-cache')
import json
from pathlib import Path
import tempfile
import unittest
import zipfile

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from lxml import etree

from trw.cli import resolve
from trw.engine import render, paragraphs

EXAMPLE = Path(__file__).resolve().parents[1] / 'examples/hello_world'


class EngineTests(unittest.TestCase):
    def test_split_tags_headers_and_layout(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            document = Document()
            document.sections[0].left_margin = 777240
            p = document.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run('Before ').italic = True
            p.add_run('@#$_na').bold = True
            p.add_run('me')
            p.add_run(' after').italic = True
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            document.sections[0].header.paragraphs[0].text = 'Report @#$_name'
            document.sections[0].footer.paragraphs[0].text = 'Footer stays'
            document.add_table(rows=1, cols=1).cell(0, 0).text = '@#$_name'
            document.save(base / 'template.docx')
            render(base / 'template.docx', base / 'out.docx',
                   {'@#$_name': {'type': 'text', 'value': 'A'}})
            out = Document(base / 'out.docx')
            self.assertEqual(out.paragraphs[0].text, 'Before A after')
            self.assertTrue(out.paragraphs[0].runs[0].italic)
            self.assertTrue(out.paragraphs[0].runs[1].bold)
            self.assertTrue(out.paragraphs[0].runs[-1].italic)
            self.assertTrue(out.paragraphs[0].paragraph_format.keep_with_next)
            self.assertEqual(out.sections[0].left_margin, 777240)
            self.assertEqual(out.sections[0].header.paragraphs[0].text, 'Report A')
            self.assertEqual(out.tables[0].cell(0, 0).text, 'A')
            self.assertEqual(len(out.element.xpath('.//w:br[@w:type="page"]')), 1)

    def test_no_silent_missing_tags_or_overwrites(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            d = Document()
            d.add_paragraph('@#$_missing')
            d.save(base / 'template.docx')
            with self.assertRaisesRegex(ValueError, 'No mapping'):
                render(base / 'template.docx', base / 'out.docx', {})
            self.assertFalse((base / 'out.docx').exists())
            (base / 'out.docx').write_bytes(b'keep this')
            with self.assertRaises(FileExistsError):
                render(base / 'template.docx', base / 'out.docx',
                       {'@#$_missing': {'type': 'text', 'value': 'ok'}})
            self.assertEqual((base / 'out.docx').read_bytes(), b'keep this')

    def test_longer_tag_not_partially_replaced(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            d = Document()
            d.add_paragraph('@#$_name @#$_name_long @#$_name')
            d.save(base / 'template.docx')
            render(base / 'template.docx', base / 'out.docx', {
                '@#$_name': {'type': 'text', 'value': 'A'},
                '@#$_name_long': {'type': 'text', 'value': 'B'}})
            self.assertEqual(Document(base / 'out.docx').paragraphs[0].text, 'A B A')

    def test_supplied_example_both_companies(self):
        config = json.loads((EXAMPLE / 'report.json').read_text())
        template = Document(EXAMPLE / config['template'])
        original_sections = [etree.tostring(s._sectPr) for s in template.sections]
        original_breaks = len(template.element.xpath('.//w:br[@w:type="page"]'))
        with tempfile.TemporaryDirectory() as folder:
            for report_id, total, male, female in [('A', 30, 17, 13), ('B', 3, 0, 3)]:
                target = Path(folder) / f'{report_id}.docx'
                status = render(EXAMPLE / config['template'], target,
                                resolve(config, EXAMPLE, report_id))
                self.assertEqual(len(status['tags_populated']), 6)
                output = Document(target)
                text = '\n'.join(p.text for p in paragraphs(output))
                self.assertNotIn('@#$_', text)
                self.assertIn(f'the {total} employees', text)
                self.assertIn(f'There are {male} males in company {report_id}', text)
                self.assertIn(f'There are {female} females in company {report_id}', text)
                data_table = output.tables[1]
                self.assertEqual([c.text for c in data_table.rows[0].cells], ['Age', 'Count'])
                self.assertEqual(sum(int(r.cells[1].text) for r in data_table.rows[1:]), total)
                self.assertEqual(len(output.inline_shapes), 2)
                for shape in output.inline_shapes:
                    self.assertTrue(shape._inline.docPr.get('descr'))
                self.assertEqual([etree.tostring(s._sectPr) for s in output.sections], original_sections)
                self.assertEqual(len(output.element.xpath('.//w:br[@w:type="page"]')), original_breaks)
                with zipfile.ZipFile(target) as archive:
                    self.assertIsNone(archive.testzip())


if __name__ == '__main__':
    unittest.main()
