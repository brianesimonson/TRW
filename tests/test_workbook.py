import os
os.environ.setdefault('MPLCONFIGDIR','/tmp/trw-matplotlib')
os.environ.setdefault('XDG_CACHE_HOME','/tmp/trw-cache')
from datetime import date
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn
import openpyxl

from trw.formulas import Calculator,FormulaError,format_value
from trw.workbook import ReportWorkbook,parse_formats,read_source,WorkbookError
from trw.pipeline import run_workbook
from trw.document import paragraphs,text_of,generate,DocumentError

ROOT=Path(__file__).resolve().parents[1]
ADVANCED=ROOT/'examples/legacy/advanced/Demo Project 1.xlsm'
FINISHED=ROOT/'examples/legacy/finished/The Report Writer - Finished Example.xlsm'


class FormulaTests(unittest.TestCase):
    def test_all_supplied_formula_cells_without_cached_values(self):
        count=0
        for path in (ROOT/'examples/legacy').rglob('*.xlsm'):
            if path.parent.name=='advanced-corrected':continue
            book=openpyxl.load_workbook(path)
            calc=Calculator(book,today=date(2026,10,13))
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        if cell.data_type=='f':
                            calc.cell(sheet.title,cell.coordinate)
                            count+=1
            self.assertEqual(calc.cell('Date','D2'),2026)
            book.close()
        self.assertEqual(count,641)

    def test_lazy_if_names_and_invalidation(self):
        book=openpyxl.Workbook();sheet=book.active;sheet.title='Data'
        sheet['A1']=2;sheet['A2']=3;sheet['B1']='=SUM(A1:A2)';sheet['B2']='=IF(A1=2,B1,1/0)'
        calc=Calculator(book)
        self.assertEqual(calc.cell('Data','B2'),5)
        sheet['A2']=7;calc.invalidate()
        self.assertEqual(calc.cell('Data','B2'),9)

    def test_unsupported_and_circular_fail_at_cell(self):
        book=openpyxl.Workbook();s=book.active;s['A1']='=UNSUPPORTED(1)'
        with self.assertRaisesRegex(FormulaError,'Sheet!A1.*Unsupported Excel function'):
            Calculator(book).cell('Sheet','A1')
        s['A1']='=B1';s['B1']='=A1'
        with self.assertRaisesRegex(FormulaError,'Circular'):
            Calculator(book).cell('Sheet','A1')

    def test_excel_formats(self):
        self.assertEqual(format_value(1234,'#,##0'),'1,234')
        self.assertEqual(format_value(.566666667,'0.0%'),'56.7%')
        self.assertEqual(format_value(2.345,'0.00'),'2.35')
        self.assertEqual(format_value(-42,'$#,##0;($#,##0)'), '($42)')
        self.assertEqual(parse_formats('format1Generalformat2#,##0',2),['General','#,##0'])


class WorkbookTests(unittest.TestCase):
    def test_id_change_imports_ranges_sections_and_index(self):
        workbook=ReportWorkbook(ADVANCED,today=date(2026,10,13))
        try:
            workbook.prepare('A')
            self.assertEqual(workbook.value('Report Numbers','C4'),30)
            self.assertTrue(workbook.sections()[1]['include'])
            self.assertEqual(len(workbook.imported),4)
            self.assertIn('CompA_Table1.csv',workbook.value('index_temp','A1'))
            workbook.prepare('B')
            self.assertEqual(workbook.value('Report Numbers','C4'),3)
            self.assertFalse(workbook.sections()[1]['include'])
            self.assertIn('female',workbook.sections()[1]['reason'])
            self.assertIn('CompB_Table1.csv',workbook.value('index_temp','A1'))
            self.assertEqual(workbook.mass_elements()['@#$_tablem01']['values'][1],['30-50','2'])
            self.assertEqual(workbook.value('Figure 1','E4'),None) # cleared stale A's third row
            self.assertIn('vba_bypassed',workbook.events)
        finally:
            workbook.close()

    def test_python_macro_extension(self):
        w=ReportWorkbook(ADVANCED,macro_hook=lambda w:setattr(w.sheet('Table 1')['F2'],'value',100))
        try:
            w.prepare('A')
            self.assertEqual(w.value('Report Numbers','C4'),126)
            self.assertIn('python_macro_hook',w.events)
            self.assertNotIn('vba_bypassed',w.events)
        finally:w.close()

    def test_legacy_bullet_variants_and_empty_fallbacks(self):
        for mode,phrase in [(1,'no common causes'),(2,'Medical Review'),(3,'Data Processing')]:
            w=ReportWorkbook(ADVANCED)
            try:
                sheet=w.sheet('Bullet 1');sheet['F2']=None;sheet['E3']='Detail';sheet['F3']=0;sheet['G3']=2
                sheet['E4']='Heading';sheet['F4']=1
                w.calc.invalidate();w.transform_bullets('Bullet 1',mode)
                self.assertIn(phrase,w.value('Bullet 1','B2'))
                self.assertEqual(sheet['A3'].font.name,'Wingdings')
                self.assertEqual(w.value('Bullet 1','B3'),'Detail (2)' if mode==1 else 'Detail')
                self.assertEqual(w.value('Bullet 1','A4'),'Heading')
                self.assertIsNone(w.value('Bullet 1','B4'))
                self.assertEqual(sheet['A4'].font.name,'Times New Roman')
            finally:w.close()

    def test_import_offsets_blank_clearing_and_quoted_csv(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);(base/'Results').mkdir()
            (base/'Results/data.csv').write_text('label,value\n"hello, world",7\n"two\nlines",9\n')
            book=openpyxl.Workbook();book.active.title='Current Parameters'
            imp=book.create_sheet('FigureDataImport');target=book.create_sheet('Target')
            imp['A2']="='Current Parameters'!A2&\"\\Results\""
            for col,val in {'E':'x','F':'Target','G':'C5','H':4,'I':2,'J':'data.csv','K':'A2'}.items():imp[f'{col}4']=val
            target['C8']='stale';book.save(base/'w.xlsx')
            w=ReportWorkbook(base/'w.xlsx');w.import_data()
            self.assertEqual(w.value('Target','C5'),'hello, world')
            self.assertEqual(w.value('Target','D6'),9)
            self.assertEqual(w.value('Target','C6'),'two\nlines')
            self.assertIsNone(w.value('Target','C8'))
            w.close()

    def test_xlsx_source_recalculates_instead_of_using_cache(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'data.xlsx';book=openpyxl.Workbook();book.active['A1']=2;book.active['B1']='=A1*3';book.save(path)
            self.assertEqual(read_source(path),[[2,6]])

    def test_import_only_workbook_has_fresh_caches_and_preserves_vba(self):
        import zipfile
        with tempfile.TemporaryDirectory() as folder:
            result=run_workbook(ADVANCED,output=folder,report_ids=['A'],import_only=True)[0]
            computed=openpyxl.load_workbook(result['output'],data_only=True)
            self.assertEqual(computed['Report Numbers']['C4'].value,30)
            self.assertEqual(computed['Current Parameters']['A11'].value,'A')
            computed.close()
            with zipfile.ZipFile(ADVANCED) as original,zipfile.ZipFile(result['output']) as prepared:
                self.assertEqual(original.read('xl/vbaProject.bin'),prepared.read('xl/vbaProject.bin'))

    def test_advanced_document_pipeline(self):
        original=ADVANCED.read_bytes()
        with tempfile.TemporaryDirectory() as folder:
            results=run_workbook(ADVANCED,output=folder,today=date(2026,10,13),skip_fields=True)
            self.assertEqual(len(results),2)
            for result in results:
                doc=Document(result['output']);text='\n'.join(text_of(p) for p in paragraphs(doc))
                self.assertNotIn('@#$_',text);self.assertNotIn('Xsection',text);self.assertNotIn('TableX',text)
                self.assertIn('2026 Employee Report',text)
                self.assertEqual(len(doc.tables),6 if result['report']=='A' else 5) # B removes the bullet table
                tables=doc.tables[1:5]
                total=30 if result['report']=='A' else 3
                for table in tables:
                    numbers=[int(cell.text) for row in table.rows for cell in row.cells if cell.text.isdigit()]
                    self.assertEqual(sum(numbers),total)
                self.assertEqual(doc.sections[0].header.paragraphs[0].text,f'Company {result["report"]} Report')
                self.assertEqual(len(doc.inline_shapes),2 if result['report']=='A' else 1)
                self.assertEqual(doc.settings.element.find(qn('w:updateFields')).get(qn('w:val')),'false')
                if result['report']=='B':
                    self.assertIn('All employees at Company B are female',text)
                    self.assertNotIn('Employees by Gender\nFigure',text)
                else:
                    self.assertTrue(any(e.get('action')=='inserted' for e in result['document_events']))
        self.assertEqual(ADVANCED.read_bytes(),original)

    def test_simple_workbook_preserves_its_actual_fixed_input_configuration(self):
        with tempfile.TemporaryDirectory() as folder:
            results=run_workbook(FINISHED,output=folder,skip_fields=True)
            for result in results:
                text='\n'.join(text_of(p) for p in paragraphs(Document(result['output'])))
                self.assertIn('the 30 employees',text)
                self.assertTrue(any('TOC' in warning for warning in result['warnings']))

    def test_corrected_copy_native_gender_merges_and_company_inputs(self):
        corrected=ROOT/'examples/legacy/advanced-corrected/Demo Project 1.xlsm'
        with tempfile.TemporaryDirectory() as folder:
            results=run_workbook(corrected,output=folder,skip_fields=True)
            for result in results:
                doc=Document(result['output'])
                merged=next(table for table in doc.tables if len(table.columns)==4)
                self.assertIs(merged.cell(0,0)._tc,merged.cell(1,1)._tc)
                self.assertIs(merged.cell(0,2)._tc,merged.cell(0,3)._tc)
                self.assertIs(merged.cell(2,0)._tc,merged.cell(3,0)._tc)
                self.assertEqual(merged.cell(2,2).text,'17' if result['report']=='A' else '0')
                self.assertEqual(merged.cell(3,2).text,'13' if result['report']=='A' else '3')
                if result['report']=='A':
                    self.assertIn('13 females in company A','\n'.join(text_of(p) for p in paragraphs(doc)))

    def test_mass_table_all_modes_scalar_formats_and_variable_stop(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);(base/'Results').mkdir()
            (base/'Results/data.csv').write_text('Name,Value\nA,0.125\nB,0.5\n,\nIgnored,0.9\n')
            book=openpyxl.Workbook();book.active.title='Current Parameters';book.active['B29']='On'
            index=book.create_sheet('index_temp')
            for row,(tag,mode,address) in enumerate([('fixed1',1,'A2:B3'),('fixed3',3,'A2:B3'),
                                                    ('variable',2,'A2:B2'),('scalar',1,'B2')],1):
                for col,value in {'A':'Results/data.csv','B':address,'C':tag,'E':str(mode),
                                  'F':'format1@format20.0%' if ':' in address else 'format10.0%'}.items():
                    index[f'{col}{row}']=value
            book.save(base/'w.xlsx');w=ReportWorkbook(base/'w.xlsx')
            try:elements=w.mass_elements()
            finally:w.close()
            self.assertEqual(elements['@#$_scalar']['value'],'12.5%')
            for tag in ['fixed1','fixed3','variable']:
                self.assertEqual(elements['@#$_'+tag]['values'],[['A','12.5%'],['B','50.0%']])
                self.assertEqual(elements['@#$_'+tag]['variable_rows'],tag=='variable')


if __name__=='__main__':
    unittest.main()
