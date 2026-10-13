from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH

from trw.document import generate,paragraphs,text_of,replace_literal,DocumentError


class DocumentFeatureTests(unittest.TestCase):
    def test_section_removes_intervening_table_and_preserves_end_suffix(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);d=Document();d.add_paragraph('Prefix Xsection1S')
            d.add_table(rows=1,cols=1).cell(0,0).text='remove this table'
            d.add_paragraph('Xsection1E suffix');d.add_paragraph('retain')
            d.save(base/'t.docx')
            generate(base/'t.docx',base/'o.docx',{},sections=[{'start':'Xsection1S','end':'Xsection1E','include':False,'reason':'Reason','bullet_mode':0}])
            out=Document(base/'o.docx');self.assertEqual(len(out.tables),0)
            self.assertEqual([p.text for p in out.paragraphs],['Prefix Reason',' suffix','retain'])

    def test_case_insensitive_split_tag_and_arbitrary_replacement(self):
        d=Document();p=d.add_paragraph();p.add_run('Prefix ').bold=True;p.add_run('@#$_Co');p.add_run('MPANY suffix')
        self.assertEqual(replace_literal(p,'@#$_company','Company B',tag=True),1)
        self.assertEqual(p.text,'Prefix Company B suffix');self.assertTrue(p.runs[0].bold)
        replace_literal(p,'SUFFIX','changed');self.assertEqual(p.text,'Prefix Company B changed')

    def test_fixed_shell_fails_instead_of_truncating(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);d=Document();t=d.add_table(rows=2,cols=2);t.cell(1,0).text='@#$_table'
            d.save(base/'t.docx')
            with self.assertRaisesRegex(DocumentError,'Fixed table shell'):
                generate(base/'t.docx',base/'o.docx',{'@#$_table':{'type':'legacy_table','mode':'shell','values':[['a','b'],['c','d']]}})
            self.assertFalse((base/'o.docx').exists())

    def test_native_table_merges_styles_headers_padding(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);d=Document();d.add_paragraph('@#$_table');d.save(base/'t.docx')
            styles=[[{'bold':True,'fill':'FFFF00','font_size':12,'horizontal':'center'},{}],[{},{}]]
            generate(base/'t.docx',base/'o.docx',{'@#$_table':{'type':'legacy_table','mode':'new',
                'values':[['Heading',''],['A','3']],'merges':[[0,0,0,1]],'cell_styles':styles,
                'header_rows':1,'widths_inches':[1,2],'padding_px':10,'padding_sides':'LR','title':'Title','alt_text':'Description'}})
            t=Document(base/'o.docx').tables[0]
            self.assertIs(t.cell(0,0)._tc,t.cell(0,1)._tc)
            self.assertTrue(t.cell(0,0).paragraphs[0].runs[0].bold)
            self.assertEqual(t.cell(0,0)._tc.tcPr.find(qn('w:shd')).get(qn('w:fill')),'FFFF00')
            self.assertIsNotNone(t.rows[0]._tr.trPr.find(qn('w:tblHeader')))
            self.assertEqual(t._tbl.tblPr.find(qn('w:tblCellMar')).find(qn('w:left')).get(qn('w:w')),'150')

    def test_missing_added_document_message_and_unknown_tag_failure(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);d=Document();d.add_paragraph('@#$_section');d.save(base/'t.docx')
            generate(base/'t.docx',base/'o.docx',{},postprocess={'documents':[{'tag':'@#$_section','filename':'missing.docx','path':str(base/'missing.docx')}]})
            self.assertIn('Section missing: File missing.docx does not exist.',Document(base/'o.docx').paragraphs[0].text)
            with self.assertRaisesRegex(DocumentError,'Unresolved template'):
                generate(base/'t.docx',base/'bad.docx',{})

    def test_five_section_headers_and_sequential_postprocessing(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);doc=Document()
            doc.add_paragraph('TableX 01 / FigureX 02 / -3,149 / Original')
            for _ in range(4):doc.add_section(WD_SECTION_START.NEW_PAGE)
            doc.save(base/'t.docx')
            specs=[{'section':i,'header':f'Header {i}','footer':f'Footer {i}',
                    'header_align':'Right','footer_align':'Center'} for i in range(1,6)]
            generate(base/'t.docx',base/'o.docx',{},postprocess={'headers':specs,
                'replacements':[('TableX 01','Table 3'),('FigureX 02','Figure 2'),
                                ('-3,149','N/A'),('Original','Updated')]})
            output=Document(base/'o.docx')
            self.assertEqual(output.paragraphs[0].text,'Table 3 / Figure 2 / N/A / Updated')
            for i,section in enumerate(output.sections,1):
                self.assertEqual(section.header.paragraphs[0].text,f'Header {i}')
                self.assertEqual(section.footer.paragraphs[0].text,f'Footer {i}')
                self.assertEqual(section.header.paragraphs[0].alignment,WD_ALIGN_PARAGRAPH.RIGHT)
                self.assertEqual(section.footer.paragraphs[0].alignment,WD_ALIGN_PARAGRAPH.CENTER)

    def test_added_document_preserves_image_relationships_and_bold_text(self):
        from trw.engine import chart_image
        from docx.shared import Inches
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);master=Document();master.add_paragraph('Before @#$_added After');master.save(base/'t.docx')
            added=Document();added.add_paragraph().add_run('Added styled content').bold=True
            added.add_picture(chart_image({'labels':['A','B'],'values':[2,3]}),width=Inches(2))
            added.save(base/'added.docx')
            generate(base/'t.docx',base/'o.docx',{},postprocess={'documents':[
                {'tag':'@#$_added','filename':'added.docx','path':str(base/'added.docx')}]})
            output=Document(base/'o.docx')
            self.assertEqual([p.text for p in output.paragraphs if p.text],['Before ','Added styled content',' After'])
            self.assertTrue(output.paragraphs[1].runs[0].bold)
            self.assertEqual(len(output.inline_shapes),1)
            image_id=output.inline_shapes[0]._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
            self.assertTrue(output.part.related_parts[image_id].blob.startswith(b'\x89PNG'))

    def test_wide_chart_fits_page_and_scales_height(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);doc=Document();doc.add_paragraph('@#$_chart');doc.save(base/'t.docx')
            generate(base/'t.docx',base/'o.docx',{'@#$_chart':{'type':'chart','labels':['A'],
                'values':[1],'width_inches':12,'height_inches':6}})
            output=Document(base/'o.docx');shape=output.inline_shapes[0];section=output.sections[0]
            self.assertLessEqual(shape.width,section.page_width-section.left_margin-section.right_margin)
            self.assertAlmostEqual(shape.height/shape.width,0.5,places=5)


if __name__=='__main__':
    unittest.main()
