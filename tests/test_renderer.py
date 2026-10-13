from pathlib import Path
import shutil
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from trw.renderer import LibreOfficeRenderer,RendererError
from trw.toc import pdf_outline,TocError
from trw.document import text_of,paragraphs

ROOT=Path(__file__).resolve().parents[1]
TEMPLATE=ROOT/'examples/hello_world/Example Project_Template.docx'


@unittest.skipUnless(shutil.which('soffice') or shutil.which('libreoffice'),'LibreOffice not installed; renderer integration checks not run')
class RendererTests(unittest.TestCase):
    def test_toc_uses_final_layout_and_changed_heading_not_old_cache(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);doc=Document(TEMPLATE)
            for p in doc.paragraphs:
                if p.text=='Conclusion':
                    p.text='Updated Conclusion'
                    p.paragraph_format.page_break_before=True
            # Add sufficient real content to change final pages, then preserve template TOC.
            for _ in range(25):doc.add_paragraph('Extra report material. '*20)
            doc.save(base/'input.docx')
            renderer=LibreOfficeRenderer()
            result=renderer.finalize(base/'input.docx',base/'out.docx',update_fields=True)
            self.assertEqual(result['indexes_updated'],1)
            self.assertGreaterEqual(result['layout_passes'],2)
            output=Document(base/'out.docx')
            self.assertEqual(output.settings.element.find(qn('w:updateFields')).get(qn('w:val')),'false')
            text='\n'.join(text_of(p) for p in paragraphs(output))
            self.assertGreaterEqual(text.count('Updated Conclusion'),2) # refreshed TOC and body
            checked=renderer.export(base/'out.docx',base/'verification')
            outline,pages=pdf_outline(checked)
            self.assertEqual(outline,result['toc_entries'])
            self.assertGreater(pages,4)
            # Cached PAGEREF text is the page label actually exported, not the old "3".
            content=output.element.xpath('.//w:sdtContent')[0]
            numbers=[]
            for p in content.iter(qn('w:p')):
                nodes=list(p.iter(qn('w:instrText')))
                if any(n.text and 'PAGEREF' in n.text for n in nodes):
                    numbers.append(list(p.iter(qn('w:t')))[-1].text)
            self.assertEqual(numbers,[e['page_label'] for e in outline])

    def test_pdf_and_legacy_doc_export(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);renderer=LibreOfficeRenderer()
            pdf=renderer.finalize(TEMPLATE,base/'out.pdf',update_fields=True,extension='.pdf')
            outline,pages=pdf_outline(base/'out.pdf')
            self.assertEqual(outline,pdf['toc_entries'])
            self.assertEqual(pages,pdf['page_count'])
            renderer.finalize(TEMPLATE,base/'out.doc',update_fields=False,extension='.doc')
            self.assertEqual((base/'out.doc').read_bytes()[:8],bytes.fromhex('D0CF11E0A1B11AE1'))

    def test_unsupported_toc_switch_is_reported_without_output(self):
        with tempfile.TemporaryDirectory() as folder:
            base=Path(folder);doc=Document(TEMPLATE)
            for node in doc.element.iter(qn('w:instrText')):
                if node.text and 'TOC ' in node.text:node.text+=' \\f '
            doc.save(base/'input.docx')
            with self.assertRaisesRegex(TocError,'Unsupported TOC switches'):
                LibreOfficeRenderer().finalize(base/'input.docx',base/'out.docx')
            self.assertFalse((base/'out.docx').exists())


if __name__=='__main__':
    unittest.main()
