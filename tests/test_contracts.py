from datetime import date
from pathlib import Path
import tempfile
import unittest

import openpyxl
from trw.workbook import ReportWorkbook, WorkbookError, parse_formats
from trw.files import wait_for_file,save_bytes
from trw.legacy import column_number,column_label,range_moves

ROOT=Path(__file__).resolve().parents[1]
ADVANCED=ROOT/'examples/legacy/advanced/Demo Project 1.xlsm'

class ContractTests(unittest.TestCase):
    def test_report_ids_stop_at_first_blank(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/'w.xlsx';b=openpyxl.Workbook();b.active.title='Current Parameters'
            s=b.create_sheet('Report Loop List');s['A2']='A';s['A3']='B';s['A5']='Ignored';b.save(p)
            w=ReportWorkbook(p)
            try:self.assertEqual(w.report_ids(),['A','B'])
            finally:w.close()

    def test_address_helpers_valid_ranges_and_explicit_invalid_inputs(self):
        self.assertEqual(column_number('AZ'),52);self.assertEqual(column_label(52),'AZ')
        self.assertEqual(range_moves('$AZ$7:$BC$12'),(52,7,3,5))
        self.assertEqual(range_moves('$B$9'),(2,9,0,0))
        with self.assertRaises(ValueError):column_number('??')
        with self.assertRaises(ValueError):column_label(0)

    def test_column_formats_and_invalid_definitions(self):
        self.assertEqual(parse_formats('format1@format2#,##0.0format30.0%',3),['@','#,##0.0','0.0%'])
        for spec in ['format2@format10','format1@','none']:
            with self.assertRaises(WorkbookError):parse_formats(spec,2)

    def test_all_postprocess_mappings_and_option_off(self):
        w=ReportWorkbook(ADVANCED,today=date(2026,10,13))
        try:
            w.prepare('A');post=w.postprocess()
            pairs=post['replacements']
            for sentinel in ('-3,149','-$3,149','-314,900.0%','-21,700.0%'):
                self.assertIn((sentinel,'N/A'),pairs)
            sheet=w.sheet('GlobalReplace');sheet['A2']='First';sheet['B2']='Second'
            sheet['A3']=None;sheet['A4']='Ignored';sheet['B4']='No'
            w.sheet('Current Parameters')['B37']='On';w.calc.invalidate()
            self.assertIn(('First','Second'),w.postprocess()['replacements'])
            self.assertNotIn(('Ignored','No'),w.postprocess()['replacements'])
            w.sheet('Current Parameters')['B37']='Off';w.calc.invalidate()
            self.assertNotIn(('First','Second'),w.postprocess()['replacements'])
            self.assertTrue(any(a.startswith('TableX ') for a,b in pairs))
            self.assertTrue(any(a.startswith('FigureX ') for a,b in pairs))
        finally:w.close()

    def test_owned_file_readiness_timeout_and_nonoverwrite(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/'file';save_bytes(b'original',p)
            self.assertEqual(wait_for_file(p),p)
            with self.assertRaises(FileExistsError):save_bytes(b'replaced',p)
            self.assertEqual(p.read_bytes(),b'original')
            with self.assertRaises(FileNotFoundError):wait_for_file(Path(folder)/'missing',timeout=0.01,interval=0.01)

if __name__=='__main__':unittest.main()
