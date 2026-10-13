"""Render an existing Word template from resolved report elements.

This module accepts data, not Excel formulas. Adapters can supply those data from
CSV, a workbook, or a future Excel ribbon integration.
"""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

TAG = re.compile(r"@#\$_[A-Za-z0-9_]+")


def paragraphs(document):
    """Include nested tables, headers, footers, and text boxes; deduplicate parts."""
    roots = [(document.element, document)]
    seen = set()
    for section in document.sections:
        for story in (section.header, section.footer, section.first_page_header,
                      section.first_page_footer, section.even_page_header,
                      section.even_page_footer):
            if story._has_definition and story.part.partname not in seen:
                roots.append((story._element, story))
                seen.add(story.part.partname)
    for root, owner in roots:
        for node in list(root.iter(qn("w:p"))):
            yield Paragraph(node, owner)


def replace_text(paragraph, tag, value):
    """Replace across runs without discarding paragraph or unrelated run styles."""
    runs = paragraph.runs
    text = ''.join(r.text for r in runs)
    for match in reversed(list(re.finditer(re.escape(tag) + r'(?![A-Za-z0-9_])', text))):
        offset = 0
        for run in runs:
            end = offset + len(run.text)
            if offset < match.end() and end > match.start():
                left = max(0, match.start() - offset)
                right = min(len(run.text), match.end() - offset)
                insert = value if offset <= match.start() < end else ''
                run.text = run.text[:left] + insert + run.text[right:]
            offset = end


def chart_image(element):
    # Figure/Agg avoids global pyplot state and does not need a display server.
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    labels, values = element['labels'], element['values']
    if len(labels) != len(values) or not labels:
        raise ValueError('Charts require equally sized, nonempty labels and values')
    width = float(element.get('width_inches', 5.5))
    height = float(element.get('height_inches', 3.0))
    fig = Figure(figsize=(width, height), layout='constrained')
    FigureCanvasAgg(fig)
    ax = fig.subplots()
    ax.bar(labels, values, color=element.get('color', '#4472C4'))
    ax.set_title(element.get('title', ''))
    ax.set_ylabel(element.get('ylabel', ''))
    if element.get('percent'):
        from matplotlib.ticker import PercentFormatter
        ax.yaxis.set_major_formatter(PercentFormatter(1))
    ax.spines[['top', 'right']].set_visible(False)
    buf = BytesIO()
    fig.savefig(buf, format='png', dpi=180)
    buf.seek(0)
    return buf


def insert_table(document, paragraph, element):
    headers, rows = element['headers'], element['rows']
    if not headers or any(len(row) != len(headers) for row in rows):
        raise ValueError('Every table row must match the header column count')
    table = document.add_table(rows=1, cols=len(headers))
    table.style = element.get('style', 'Table Grid')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = str(value)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = str(value)
    widths = element.get('column_widths_inches')
    if widths:
        if len(widths) != len(headers):
            raise ValueError('Table widths must match the column count')
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    props = table._tbl.tblPr
    caption = OxmlElement('w:tblCaption')
    caption.set(qn('w:val'), element.get('title', ''))
    props.append(caption)
    description = OxmlElement('w:tblDescription')
    description.set(qn('w:val'), element.get('alt_text', ''))
    props.append(description)
    # Retain the anchor paragraph and its page-break/spacing settings before table.
    replace_text(paragraph, element['tag'], '')
    paragraph._p.addnext(table._tbl)


def bullet_numbering(document):
    """Create real Word bullets even if the template has no List Bullet style."""
    root = document.part.numbering_part.element
    existing = [int(n.get(qn('w:abstractNumId'))) for n in root.findall(qn('w:abstractNum'))]
    abstract_id = max(existing, default=-1) + 1
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(abstract_id))
    level = OxmlElement('w:lvl')
    level.set(qn('w:ilvl'), '0')
    for name, value in [('start', '1'), ('numFmt', 'bullet'), ('lvlText', '•'), ('lvlJc', 'left')]:
        item = OxmlElement('w:' + name)
        item.set(qn('w:val'), value)
        level.append(item)
    abstract.append(level)
    root.append(abstract)
    return root.add_num(abstract_id).numId


def render(template, destination, elements):
    """Render atomically; fail rather than silently leave unresolved placeholders."""
    document = Document(template)
    all_paragraphs = list(paragraphs(document))
    available = {tag for p in all_paragraphs for tag in TAG.findall(p.text)}
    unknown = available - elements.keys()
    if unknown:
        raise ValueError('No mapping for template tags: ' + ', '.join(sorted(unknown)))
    for paragraph in all_paragraphs:
        for tag in set(TAG.findall(paragraph.text)):
            element = elements[tag]
            kind = element['type']
            if kind == 'text':
                replace_text(paragraph, tag, str(element['value']))
            elif kind == 'chart':
                if paragraph.text.strip() != tag:
                    raise ValueError(f'Chart tag must occupy its own paragraph: {tag}')
                replace_text(paragraph, tag, '')
                picture = paragraph.add_run().add_picture(
                    chart_image(element), width=Inches(element.get('width_inches', 5.5)))
                picture._inline.docPr.set('descr', element.get('alt_text', ''))
            elif kind == 'table':
                if paragraph.text.strip() != tag:
                    raise ValueError(f'Table tag must occupy its own paragraph: {tag}')
                insert_table(document, paragraph, dict(element, tag=tag))
            elif kind == 'bullets':
                if paragraph.text.strip() != tag:
                    raise ValueError(f'Bullet tag must occupy its own paragraph: {tag}')
                replace_text(paragraph, tag, '')
                anchor = paragraph._p
                numbering_id = bullet_numbering(document)
                for value in element['items']:
                    new = deepcopy(paragraph._p)
                    anchor.addnext(new)
                    p = Paragraph(new, paragraph._parent)
                    if 'style' in element:
                        p.style = element['style']
                    num = p._p.get_or_add_pPr().get_or_add_numPr()
                    num.get_or_add_ilvl().val = 0
                    num.get_or_add_numId().val = numbering_id
                    p.paragraph_format.left_indent = Inches(0.25)
                    p.paragraph_format.first_line_indent = Inches(-0.15)
                    # A copied page-break-before belongs only to the original anchor.
                    p.paragraph_format.page_break_before = False
                    p.add_run(str(value))
                    anchor = new
            else:
                raise ValueError(f'Unsupported element type: {kind}')
    remaining = {tag for p in paragraphs(document) for tag in TAG.findall(p.text)}
    if remaining:
        raise ValueError('Unresolved tags: ' + ', '.join(sorted(remaining)))
    # Word may refresh fields on opening; this does not calculate TOC page numbers.
    update = document.settings.element.find(qn('w:updateFields'))
    if update is None:
        update = OxmlElement('w:updateFields')
        document.settings.element.append(update)
    update.set(qn('w:val'), 'true')
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    import tempfile
    import os
    if destination.exists():
        raise FileExistsError(f'Refusing to overwrite {destination}')
    with tempfile.NamedTemporaryFile(dir=destination.parent, suffix='.docx', delete=False) as temp:
        temporary = Path(temp.name)
    try:
        document.save(temporary)
        # Hard link is atomic and refuses overwrite, including concurrent writers.
        os.link(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)
    return {'output': str(destination), 'tags_populated': sorted(available),
            'unused_mappings': sorted(elements.keys() - available)}
