"""Explicit Excel formula subset used by Report Writer's workbook contract.

No eval(), VBA execution, or cached Excel results. Unsupported syntax/functions
and circular references fail with the source cell attached to the error.
"""
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
import fnmatch
import math
import re

from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import to_excel, from_excel


class FormulaError(ValueError):
    pass


def excel_text(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, float):
        return format(value, '.15g')
    return str(value)


def number(value):
    if value is None or value == '':
        return 0
    if isinstance(value, (date, datetime)):
        return to_excel(datetime.combine(value, datetime.min.time()) if not isinstance(value, datetime) else value)
    try:
        return float(value)
    except (ValueError, TypeError) as exc:
        raise FormulaError(f'Expected number, got {value!r}') from exc


def format_value(value, code='General'):
    """Numeric/text/date formats needed by supplied workbooks; reject unknowns."""
    if value is None:
        return ''
    if code in ('General', '@', '', None):
        return excel_text(value)
    if isinstance(value, str):
        return value
    sections = code.split(';')
    section = sections[1] if number(value) < 0 and len(sections) > 1 else sections[0]
    section = re.sub(r'\[(?:Red|Black|Blue|Green|Yellow|White|Cyan|Magenta)\]', '', section, flags=re.I)
    if re.search(r'[ymdhis]', re.sub(r'"[^"]*"', '', section), re.I):
        dt = value if isinstance(value, (date, datetime)) else from_excel(number(value))
        if not isinstance(dt, (date, datetime)):
            raise FormulaError(f'Unsupported time-only format: {code}')
        date_tokens = {'yyyy': str(dt.year), 'yy': f'{dt.year % 100:02}',
                       'mmmm': dt.strftime('%B'), 'mmm': dt.strftime('%b'),
                       'mm': f'{dt.month:02}', 'm': str(dt.month),
                       'dddd': dt.strftime('%A'), 'ddd': dt.strftime('%a'),
                       'dd': f'{dt.day:02}', 'd': str(dt.day)}
        if not re.fullmatch(r'[ymd/\- .]+', section, flags=re.I):
            raise FormulaError(f'Unsupported date format: {code}')
        return re.sub(r'yyyy|mmmm|dddd|mmm|ddd|yy|mm|dd|m|d',
                      lambda m: date_tokens[m[0].lower()], section, flags=re.I)
    # Excel alignment/fill directives have no textual representation.
    section = re.sub(r'_.|\*.', '', section)
    section = re.sub(r'\\(.)', r'\1', section)
    literal = re.sub(r'"([^"]*)"', r'\1', section)
    match = re.search(r'[#0][#0,]*(?:\.([#0]+))?', literal)
    if not match:
        raise FormulaError(f'Unsupported number format: {code}')
    prefix, suffix = literal[:match.start()], literal[match.end():]
    if re.search(r'[A-Za-z\[\]<>]', prefix + suffix):
        raise FormulaError(f'Unsupported number format: {code}')
    percent = '%' in suffix
    numeric = number(value) * (100 if percent else 1)
    if number(value) < 0 and len(sections) > 1:
        numeric = abs(numeric)
    decimals = len(match[1] or '')
    rounded = Decimal(str(numeric)).quantize(Decimal(1).scaleb(-decimals), rounding=ROUND_HALF_UP)
    formatted = format(rounded, (',' if ',' in match[0] else '') + f'.{decimals}f')
    optional = (match[1] or '').count('#')
    if optional and '.' in formatted:
        for _ in range(optional):
            if formatted.endswith('0'):
                formatted = formatted[:-1]
        formatted = formatted.rstrip('.')
    return prefix + formatted + suffix


TOKEN = re.compile(r'\s*(?:("(?:[^"]|"")*")|(\'(?:[^\']|\'\')*\')|'
                   r'(\$?[A-Za-z]{1,3}\$?\d+)(?![A-Za-z0-9_.])|'
                   r'(\d+(?:\.\d*)?(?:[Ee][+-]?\d+)?|\.\d+)|'
                   r'([A-Za-z_][A-Za-z0-9_.]*)|(<=|>=|<>|[+\-*/^&=<>%,():!]))')
PRECEDENCE = {'=': 10, '<>': 10, '<': 10, '>': 10, '<=': 10, '>=': 10,
              '&': 20, '+': 30, '-': 30, '*': 40, '/': 40, '^': 50}


class Parser:
    def __init__(self, formula):
        self.tokens = []
        position = 0
        formula = formula.lstrip('=')
        while position < len(formula):
            if not formula[position:].strip():
                break
            match = TOKEN.match(formula, position)
            if not match:
                raise FormulaError(f'Unsupported formula syntax near {formula[position:position+30]!r}')
            kind = next(i for i, value in enumerate(match.groups()) if value is not None)
            self.tokens.append((kind, match[kind + 1]))
            position = match.end()
        self.tokens.append((5, 'EOF'))
        self.i = 0

    def peek(self):
        return self.tokens[self.i][1]

    def take(self, expected=None):
        token = self.tokens[self.i]
        if expected is not None and token[1] != expected:
            raise FormulaError(f'Expected {expected}, got {token[1]}')
        self.i += 1
        return token

    def parse(self):
        node = self.expression()
        self.take('EOF')
        return node

    def expression(self, minimum=0):
        kind, text = self.take()
        if text in ('+', '-'):
            left = ('unary', text, self.expression(49))
        elif text == '(':
            left = self.expression()
            self.take(')')
        elif kind == 0:
            left = ('literal', text[1:-1].replace('""', '"'))
        elif kind == 3:
            left = ('literal', float(text))
        elif kind in (1, 2, 4):
            sheet = None
            if self.peek() == '!':
                self.take('!')
                sheet = text[1:-1].replace("''", "'") if kind == 1 else text
                kind, text = self.take()
                if kind not in (2, 4):
                    raise FormulaError('Expected cell or named range after sheet qualifier')
            if kind == 4 and self.peek() == '(' and sheet is None:
                self.take('(')
                args = []
                if self.peek() != ')':
                    while True:
                        args.append(self.expression())
                        if self.peek() != ',':
                            break
                        self.take(',')
                self.take(')')
                left = ('call', text.upper(), args)
            elif kind == 4 and text.upper() in ('TRUE', 'FALSE') and sheet is None:
                left = ('literal', text.upper() == 'TRUE')
            elif kind == 2:
                if self.peek() == ':':
                    self.take(':')
                    other_kind, other = self.take()
                    if other_kind != 2:
                        raise FormulaError('Expected cell after range colon')
                    text += ':' + other
                left = ('reference', sheet, text)
            else:
                left = ('name', sheet, text)
        else:
            raise FormulaError(f'Unexpected formula token: {text}')
        while True:
            if self.peek() == '%':
                self.take('%')
                left = ('percent', left)
                continue
            operator = self.peek()
            precedence = PRECEDENCE.get(operator, -1)
            if precedence < minimum:
                break
            self.take()
            right = self.expression(precedence if operator == '^' else precedence + 1)
            left = ('binary', operator, left, right)
        return left


def flatten(value):
    if isinstance(value, list):
        for item in value:
            yield from flatten(item)
    else:
        yield value


def compare(left, right, operator):
    if isinstance(left, str) and isinstance(right, str):
        left, right = left.casefold(), right.casefold()
    elif left is None or right is None:
        # Empty cell equals both empty string and numeric zero in Excel.
        left = '' if left is None and isinstance(right, str) else (0 if left is None else left)
        right = '' if right is None and isinstance(left, str) else (0 if right is None else right)
    if operator == '=':
        return left == right
    if operator == '<>':
        return left != right
    # Excel sorts numbers before text in lookup/comparison operations.
    if isinstance(left, str) != isinstance(right, str):
        left = (1, left) if isinstance(left, str) else (0, left)
        right = (1, right) if isinstance(right, str) else (0, right)
    return {'<': lambda: left < right, '>': lambda: left > right,
            '<=': lambda: left <= right, '>=': lambda: left >= right}[operator]()


class Calculator:
    def __init__(self, workbook, today=None):
        self.workbook = workbook
        self.today = today or date.today()
        self.cache = {}
        self.active = set()
        self.parsed = {}

    def sheet(self, name):
        for sheet in self.workbook:
            if sheet.title.casefold() == name.casefold():
                return sheet
        raise FormulaError(f'Unknown worksheet: {name}')

    def invalidate(self):
        self.cache.clear()

    def cell(self, sheet, address):
        worksheet = self.sheet(sheet)
        address = address.replace('$', '').upper()
        key = (worksheet.title, address)
        if key in self.cache:
            return self.cache[key]
        if key in self.active:
            raise FormulaError(f'Circular formula at {sheet}!{address}')
        self.active.add(key)
        try:
            cell = worksheet[address]
            value = cell.value
            if cell.data_type == 'e':
                raise FormulaError(f'Excel error {value}')
            if cell.data_type == 'f':
                if not isinstance(value, str):
                    raise FormulaError('Unsupported array/data-table formula')
                node = self.parsed.setdefault(value, Parser(value).parse())
                value = self.evaluate(node, worksheet.title)
            self.cache[key] = value
            return value
        except (FormulaError, ArithmeticError, ValueError, TypeError, KeyError) as exc:
            raise FormulaError(f'{worksheet.title}!{address}: {exc}') from exc
        finally:
            self.active.remove(key)

    def rectangle(self, sheet, address):
        min_col, min_row, max_col, max_row = range_boundaries(address)
        if None in (min_col, min_row, max_col, max_row):
            raise FormulaError('Whole-column/row references are not supported')
        if (max_col - min_col + 1) * (max_row - min_row + 1) > 1_000_000:
            raise FormulaError('Formula range exceeds one million cells')
        from openpyxl.utils import get_column_letter
        return [[self.cell(sheet, f'{get_column_letter(c)}{r}') for c in range(min_col, max_col + 1)]
                for r in range(min_row, max_row + 1)]

    def evaluate(self, node, current_sheet):
        kind, *args = node
        if kind == 'literal':
            return args[0]
        if kind == 'reference':
            sheet, address = args
            return self.rectangle(sheet or current_sheet, address) if ':' in address else self.cell(sheet or current_sheet, address)
        if kind == 'name':
            sheet, name = args
            collection = self.sheet(sheet or current_sheet).defined_names
            definition = next((d for n, d in collection.items() if n.casefold() == name.casefold()), None)
            if definition is None:
                definition = next((d for n, d in self.workbook.defined_names.items() if n.casefold() == name.casefold()), None)
            if definition is None:
                raise FormulaError(f'Unknown named range: {name}')
            return self.evaluate(Parser(definition.attr_text).parse(), sheet or current_sheet)
        if kind == 'unary':
            return number(self.evaluate(args[1], current_sheet)) * (-1 if args[0] == '-' else 1)
        if kind == 'percent':
            return number(self.evaluate(args[0], current_sheet)) / 100
        if kind == 'binary':
            op, left, right = args
            left, right = self.evaluate(left, current_sheet), self.evaluate(right, current_sheet)
            if op == '&':
                return excel_text(left) + excel_text(right)
            if op in ('=', '<>', '<', '>', '<=', '>='):
                return compare(left, right, op)
            a, b = number(left), number(right)
            return {'+': lambda: a + b, '-': lambda: a - b, '*': lambda: a * b,
                    '/': lambda: a / b, '^': lambda: a ** b}[op]()
        if kind != 'call':
            raise FormulaError(f'Unsupported formula node: {kind}')
        func, expressions = args
        if func == 'IF':
            if len(expressions) not in (2, 3):
                raise FormulaError('IF requires two or three arguments')
            branch = 1 if self.evaluate(expressions[0], current_sheet) else 2
            return self.evaluate(expressions[branch], current_sheet) if branch < len(expressions) else False
        values = [self.evaluate(expr, current_sheet) for expr in expressions]
        flat = list(flatten(values))
        if func == 'SUM':
            return sum(number(v) for v in flat if isinstance(v, (int, float)) and not isinstance(v, bool))
        if func == 'COUNTA':
            return sum(v is not None for v in flat)
        if func == 'COUNTIF':
            criteria = values[1]
            def matches(v):
                if isinstance(criteria, str):
                    match = re.match(r'^(<=|>=|<>|=|<|>)(.*)$', criteria)
                    if match:
                        literal = match[2]
                        try:
                            literal = float(literal)
                        except ValueError:
                            pass
                        return compare(v, literal, match[1])
                    return fnmatch.fnmatchcase(excel_text(v).casefold(), criteria.casefold())
                return compare(v, criteria, '=')
            return sum(matches(v) for v in flatten(values[0]))
        if func in ('CONCATENATE', 'CONCAT'):
            return ''.join(excel_text(v) for v in flat)
        if func == 'LOWER':
            return excel_text(values[0]).lower()
        if func == 'TEXT':
            return format_value(values[0], values[1])
        if func == 'OR':
            return any(bool(v) for v in flat)
        if func == 'AND':
            return all(bool(v) for v in flat)
        if func == 'NOT':
            return not values[0]
        if func == 'TODAY':
            return to_excel(datetime.combine(self.today, datetime.min.time()), self.workbook.epoch)
        if func in ('DAY', 'MONTH', 'YEAR'):
            dt = values[0] if isinstance(values[0], (date, datetime)) else from_excel(number(values[0]), self.workbook.epoch)
            return getattr(dt, func.lower())
        if func == 'LOOKUP':
            lookup, keys, results = values
            keys, results = list(flatten(keys)), list(flatten(results))
            if len(keys) != len(results):
                raise FormulaError('LOOKUP vectors differ in length')
            candidates = [i for i, key in enumerate(keys) if key is not None and compare(key, lookup, '<=')]
            if not candidates:
                raise FormulaError('LOOKUP: #N/A')
            return results[candidates[-1]]
        if func == 'VLOOKUP':
            lookup, table, index, *mode = values
            index = int(number(index)) - 1
            if index < 0 or any(index >= len(row) for row in table):
                raise FormulaError('VLOOKUP: invalid column index')
            approximate = bool(mode[0]) if mode else True
            candidates = [row for row in table if row[0] is not None and compare(row[0], lookup, '<=' if approximate else '=')]
            if not candidates:
                raise FormulaError('VLOOKUP: #N/A')
            return candidates[-1 if approximate else 0][index]
        raise FormulaError(f'Unsupported Excel function: {func}')
